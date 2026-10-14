"""Extract words and coordinates for exact chunk highlighting, with optional OCR."""
import io
import os
from pathlib import Path
from zipfile import ZipFile


def ocr_words(image):
    import pytesseract
    from pytesseract import Output
    command = os.getenv('TESSERACT_CMD')
    if command:
        pytesseract.pytesseract.tesseract_cmd = command
    try:
        data = pytesseract.image_to_data(image, lang=os.getenv('OCR_LANG', 'eng'), output_type=Output.DICT, timeout=90)
    except pytesseract.TesseractNotFoundError:
        raise ValueError('OCR needs Tesseract. Install it and set TESSERACT_CMD in .env, or use Docker.') from None
    words, boxes = [], []
    for i, raw in enumerate(data['text']):
        text = raw.strip()
        if text:
            words.append(text)
            boxes.append([data['left'][i], data['top'][i], data['left'][i]+data['width'][i], data['top'][i]+data['height'][i]])
    return words, boxes


def extract(path, ocr, progress):
    import pymupdf as fitz
    from PIL import Image, ImageOps
    suffix = Path(path).suffix.lower()
    pages = []
    if suffix == '.pdf':
        with fitz.open(path) as pdf:
            if pdf.is_encrypted or len(pdf) > 250:
                raise ValueError('Use an unencrypted PDF with 250 pages or fewer.')
            total_words = 0
            for i, page in enumerate(pdf):
                progress('extracting', 10 + int(30 * i / max(1,len(pdf))))
                if page.rect.width * page.rect.height > 4_000_000:
                    raise ValueError('PDF page dimensions are too large.')
                native = page.get_text('words', sort=True)
                words = [str(w[4]) for w in native]
                boxes = [list(fitz.Rect(w[:4]) * page.rotation_matrix) for w in native]
                if len(' '.join(words).strip()) < 30 and ocr:
                    progress(f'OCR page {i+1}', 10 + int(30*i/max(1,len(pdf))))
                    pix = page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
                    image = Image.frombytes('RGB', [pix.width,pix.height],pix.samples)
                    words, boxes = ocr_words(image)
                    boxes = [[b[0]*page.rect.width/pix.width,b[1]*page.rect.height/pix.height,b[2]*page.rect.width/pix.width,b[3]*page.rect.height/pix.height] for b in boxes]
                total_words += len(words)
                if total_words > 300_000:
                    raise ValueError('Document exceeds 300,000 extracted words.')
                pages.append({'page':i+1,'words':words,'boxes':boxes})
    elif suffix in {'.png','.jpg','.jpeg','.tiff','.bmp'}:
        if not ocr:
            raise ValueError('Enable OCR to upload images.')
        with Image.open(path) as raw:
            if raw.width*raw.height > 20_000_000:
                raise ValueError('Image exceeds 20 megapixels.')
            image=ImageOps.exif_transpose(raw).convert('RGB')
            progress('OCR image',25)
            words,boxes=ocr_words(image)
            pages=[{'page':1,'words':words,'boxes':boxes}]
    else:
        if suffix == '.docx':
            from docx import Document
            with ZipFile(path) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50*1024*1024:
                    raise ValueError('Expanded DOCX exceeds 50 MB.')
            doc=Document(path)
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        elif suffix in {'.txt','.md'}:
            text=Path(path).read_text(encoding='utf-8-sig')
        else:
            raise ValueError('Unsupported file format.')
        if len(text)>2_000_000:
            raise ValueError('Extracted text exceeds 2 million characters.')
        pages=[{'page':1,'words':text.split(),'boxes':[]}]
    progress('chunking',45)
    passages=[]
    for page in pages:
        words=page['words']
        for start in range(0,len(words),130):
            passages.append({'page':page['page'],'text':' '.join(words[start:start+160]),'boxes':page['boxes'][start:start+160]})
            if start+160>=len(words):
                break
    if not passages:
        raise ValueError('No readable text. Enable OCR for scanned documents.')
    return passages


def preview(path, page_number, boxes):
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageOps
    suffix=Path(path).suffix.lower()
    if suffix=='.pdf':
        with fitz.open(path) as pdf:
            page=pdf[page_number-1]
            pix=page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False)
            image=Image.frombytes('RGB',[pix.width,pix.height],pix.samples).convert('RGBA')
            sx,sy=pix.width/page.rect.width,pix.height/page.rect.height
    else:
        with Image.open(path) as original:
            image=ImageOps.exif_transpose(original).convert('RGBA')
        sx=sy=1
    layer=Image.new('RGBA',image.size)
    drawing=ImageDraw.Draw(layer)
    for x0,y0,x1,y1 in boxes:
        drawing.rectangle([x0*sx,y0*sy,x1*sx,y1*sy],fill=(250,200,40,75))
    image=Image.alpha_composite(image,layer).convert('RGB')
    image.thumbnail((1600,2200))
    output=io.BytesIO();image.save(output,format='PNG')
    return output.getvalue()
